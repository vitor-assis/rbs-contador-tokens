"""Extração de texto dos arquivos da base de conhecimento."""
from __future__ import annotations

import io
import os
from dataclasses import dataclass, field
from html.parser import HTMLParser

# Aumente ao mudar o que ``extract`` devolve: a UI usa como chave de cache, para não
# reaproveitar resultados de uma versão anterior (ex.: imagens sem dimensões).
EXTRACTOR_VERSION = 2

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".xml"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".html", ".htm", ".pdf", ".docx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
# Formatos aceitos por todos os provedores; os demais precisam ser convertidos antes do envio.
PROVIDER_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


@dataclass
class ExtractedFile:
    name: str
    extension: str
    text: str = ""
    pages: int = 0
    page_sizes: list[tuple[float, float]] = field(default_factory=list)  # PDF, em pontos (1/72")
    width: int = 0   # imagem, em px
    height: int = 0
    is_image: bool = False
    error: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def chars(self) -> int:
        return len(self.text)

    @property
    def ok(self) -> bool:
        return self.error is None and not self.is_image


def decode_text(data: bytes) -> str:
    """Decodifica bytes como UTF-8 (com ou sem BOM), caindo para cp1252/latin-1."""
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


class _HTMLText(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template"}
    _BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _HTMLText()
    parser.feed(html)
    parser.close()
    lines = [" ".join(line.split()) for line in "".join(parser.parts).splitlines()]
    return "\n".join(line for line in lines if line)


def _extract_pdf(data: bytes, result: ExtractedFile) -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    result.pages = len(reader.pages)
    texts = []
    for page in reader.pages:
        box = page.cropbox
        w, h = float(box.width), float(box.height)
        if (page.rotation or 0) % 180 == 90:
            w, h = h, w
        result.page_sizes.append((w, h))
        texts.append(page.extract_text() or "")
    result.text = "\n\n".join(t.strip() for t in texts if t.strip())
    if result.pages and not result.text.strip():
        result.warnings.append(
            "Nenhum texto extraído: o PDF parece ser escaneado (só imagens). "
            "Provedores com leitura nativa de PDF ainda cobrarão as páginas como imagem."
        )


def _extract_docx(data: bytes, result: ExtractedFile) -> None:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    result.text = "\n".join(parts)


def extract(name: str, data: bytes) -> ExtractedFile:
    """Extrai o texto de um arquivo. Nunca lança exceção: erros vão em ``error``."""
    ext = os.path.splitext(name)[1].lower()
    result = ExtractedFile(name=name, extension=ext)

    if ext in IMAGE_EXTENSIONS:
        result.is_image = True
        try:
            from PIL import Image

            with Image.open(io.BytesIO(data)) as img:
                result.width, result.height = img.size
        except Exception as e:
            result.error = f"Não foi possível ler a imagem: {e}"
            return result
        if ext not in PROVIDER_IMAGE_EXTENSIONS:
            result.warnings.append("Formato não aceito pelos provedores (use PNG, JPEG, GIF ou WebP); "
                                   "estimado como se fosse convertido.")
        return result
    if ext not in SUPPORTED_EXTENSIONS:
        result.error = f"Formato '{ext or '(sem extensão)'}' não suportado."
        return result

    try:
        if ext in TEXT_EXTENSIONS:
            result.text = decode_text(data)
        elif ext in {".html", ".htm"}:
            result.text = html_to_text(decode_text(data))
            result.warnings.append("Tags HTML, scripts e estilos removidos; só o texto visível é contado.")
        elif ext == ".pdf":
            _extract_pdf(data, result)
        elif ext == ".docx":
            _extract_docx(data, result)
    except Exception as e:  # arquivo corrompido, senha, etc.
        result.error = f"Não foi possível ler o arquivo: {e}"
    return result
